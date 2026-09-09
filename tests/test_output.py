"""obsidian_gen / excel_report 输出模块测试（SPEC §15）。

使用构造的假数据（5 篇 Paper，覆盖不同 source / cluster /
有无 DOI / 各种 pdf_status / 标题冲突与特殊字符）。
"""

import os
import re

import pytest
import yaml
from openpyxl import load_workbook

from paper_agent.excel_report import HEADERS, SHEET_NAME, write_excel
from paper_agent.models import Paper, PaperAnalysis
from paper_agent.obsidian_gen import generate_vault

QUERY = "graph neural networks for drug discovery"

CLUSTER_A = "GNN-分子表示"
CLUSTER_B = "蛋白质结构预测"

# 含 Obsidian 非法字符 : " [ ]，用于验证文件名 safe 化
TITLE_1 = 'Graph Neural Networks for Drug Discovery: A Survey "v2" [Updated]'


def make_papers() -> list[Paper]:
    return [
        Paper(
            id="arxiv:2301.00001",
            source="arxiv",
            title=TITLE_1,
            authors=["Alice Zhang", "Bob Li", "Carol Wang", "Dave Chen"],
            abstract="We survey graph neural networks for drug discovery.",
            year=2024,
            venue="",
            doi=None,
            url="https://arxiv.org/abs/2301.00001",
            pdf_url="https://arxiv.org/pdf/2301.00001.pdf",
            citation_count=None,
            relevance_score=95.0,
            authority_score=60.0,
            pdf_status="ok",
            local_pdf="pdfs/paper1.pdf",
            zotero_key="ABCD1234",
        ),
        Paper(
            id="s2:CorpusId:123456",
            source="semantic_scholar",
            title="Molecular Representation Learning with Graph Transformers",
            authors=["Eve Smith"],
            abstract="Graph transformers for molecular property prediction.",
            year=None,                      # 无年份
            venue="Nature Machine Intelligence",
            doi="10.1038/s42256-023-00001-x",
            url="https://www.semanticscholar.org/paper/123456",
            pdf_url=None,
            citation_count=152,
            relevance_score=88.0,
            authority_score=92.0,
            pdf_status="paywalled",
            pdf_note="需要机构认证或付费订阅",
        ),
        Paper(
            id="doi:10.1126/science.abc1234",
            source="crossref",
            title="Highly Accurate Protein Structure Prediction with AlphaFold",
            authors=["John Jumper", "Demis Hassabis"],
            abstract="AlphaFold can predict protein structures with atomic accuracy.",
            year=2021,
            venue="Nature",
            doi="10.1126/science.abc1234",
            url="https://doi.org/10.1126/science.abc1234",
            pdf_url=None,
            citation_count=12000,
            relevance_score=70.0,
            authority_score=98.0,
            pdf_status="no_pdf",
        ),
        Paper(
            id="pmid:38123456",
            source="pubmed",
            title="Deep Learning for Protein-Ligand Docking",
            authors=["Frank Liu"],
            abstract="",                     # 空摘要
            year=2023,
            venue="J Chem Inf Model",
            doi=None,                        # 无 DOI
            url="https://pubmed.ncbi.nlm.nih.gov/38123456/",
            pdf_url=None,
            citation_count=None,
            relevance_score=65.0,
            authority_score=40.0,
            pdf_status="error",
            pdf_note="连接超时",
        ),
        Paper(
            id="arxiv:2301.00001v2",
            source="arxiv",
            title=TITLE_1,                   # 与第 1 篇同名 -> 文件名冲突消歧
            authors=["Alice Zhang", "Bob Li"],
            abstract="Updated version.",
            year=2024,
            venue="",
            doi=None,
            url="https://arxiv.org/abs/2301.00001v2",
            pdf_url="https://arxiv.org/pdf/2301.00001v2.pdf",
            citation_count=5,
            relevance_score=90.0,
            authority_score=30.0,
            pdf_status="pending",
        ),
    ]


def make_analyses(papers: list[Paper]) -> dict[str, PaperAnalysis]:
    clusters = {
        "arxiv:2301.00001": CLUSTER_A,
        "s2:CorpusId:123456": CLUSTER_A,
        "arxiv:2301.00001v2": CLUSTER_A,
        "doi:10.1126/science.abc1234": CLUSTER_B,
        "pmid:38123456": CLUSTER_B,
    }
    analyses = {}
    for p in papers:
        analyses[p.id] = PaperAnalysis(
            paper_id=p.id,
            summary=f"本文围绕{p.title[:20]}展开研究。提出了新的方法框架。实验表明性能显著提升。",
            deep_read="建议精读实验部分，因为包含与基线的详细对比。",
            skim="相关工作一节可略读，多为已有综述内容的重复。",
            questions=["该方法在小样本场景下是否仍然有效？",
                        "能否推广到其他分子任务？"],
            cluster=clusters[p.id],
        )
    return analyses


@pytest.fixture()
def papers():
    return make_papers()


@pytest.fixture()
def analyses(papers):
    return make_analyses(papers)


@pytest.fixture()
def vault_dir(tmp_path, papers, analyses):
    out = tmp_path / "obsidian_vault"
    generate_vault(QUERY, papers, analyses, str(out))
    return out


# ---------------- Obsidian ----------------

def _parse_frontmatter(text: str) -> dict:
    assert text.startswith("---\n"), "笔记缺少 frontmatter"
    end = text.index("\n---", 4)
    return yaml.safe_load(text[4:end])


_WIKILINK_RE = re.compile(r"\[\[([^\[\]]+?)\]\]")


def test_vault_structure(vault_dir, papers):
    mocs = [f for f in os.listdir(vault_dir) if f.startswith("00-MOC-")
            and f.endswith(".md")]
    assert len(mocs) == 1, f"MOC 文件不唯一: {mocs}"
    notes_dir = vault_dir / "papers"
    notes = [f for f in os.listdir(notes_dir) if f.endswith(".md")]
    assert len(notes) == len(papers), "每篇论文应对应一篇笔记"


def test_note_filenames_safe_and_unique(vault_dir):
    notes = os.listdir(vault_dir / "papers")
    assert len(notes) == len(set(notes))
    illegal = re.compile(r'[\\/:*?"<>|#^\[\]]')
    for name in notes:
        assert not illegal.search(name), f"文件名含非法字符: {name}"
        assert len(name[:-3]) <= 110  # 80 截断 + 冲突后缀余量
    # 两篇同名论文必须生成两个不同文件
    same_title = [n for n in notes if n.startswith("Graph Neural Networks for Drug Discovery")]
    assert len(same_title) == 2


def test_frontmatter_valid_yaml(vault_dir):
    for fname in os.listdir(vault_dir / "papers"):
        text = (vault_dir / "papers" / fname).read_text(encoding="utf-8")
        fm = _parse_frontmatter(text)
        assert isinstance(fm, dict), f"frontmatter 不是 YAML dict: {fname}"
        for key in ("title", "source", "year", "authors", "doi", "url",
                    "relevance", "authority", "cluster", "tags"):
            assert key in fm, f"{fname} frontmatter 缺少 {key}"
        assert isinstance(fm["authors"], list)
        assert isinstance(fm["tags"], list)
        assert "paper" in fm["tags"]


def test_no_dead_wikilinks(vault_dir):
    """解析所有 .md 中的 [[...]] 目标，逐一确认对应文件存在（零死链）。"""
    md_files = [vault_dir / f for f in os.listdir(vault_dir) if f.endswith(".md")]
    md_files += [vault_dir / "papers" / f
                 for f in os.listdir(vault_dir / "papers") if f.endswith(".md")]
    valid_targets = {f.stem for f in md_files}
    assert len(md_files) == 6  # 1 MOC + 5 notes

    all_links = 0
    for f in md_files:
        text = f.read_text(encoding="utf-8")
        for raw in _WIKILINK_RE.findall(text):
            target = raw.split("|", 1)[0].split("#", 1)[0].strip()
            all_links += 1
            assert target in valid_targets, f"{f.name} 存在死链: [[{raw}]]"
    assert all_links > 0, "应当至少存在若干 wikilink"


def test_note_sections_and_related_links(vault_dir, papers, analyses):
    notes = list((vault_dir / "papers").glob("*.md"))
    for f in notes:
        text = f.read_text(encoding="utf-8")
        for section in ("## 内容大意", "## 建议精读", "## 建议略读",
                        "## 进一步思考", "## 相关论文"):
            assert section in text, f"{f.name} 缺少小节 {section}"
        assert "返回 [[" in text and "返回目录" in text

    # 同 cluster 笔记应互相链接
    cluster_a_notes = [
        f for f in notes
        if _parse_frontmatter(f.read_text(encoding="utf-8"))["cluster"] == CLUSTER_A
    ]
    assert len(cluster_a_notes) == 3
    for f in cluster_a_notes:
        text = f.read_text(encoding="utf-8")
        others = {o.stem for o in cluster_a_notes if o != f}
        linked = {m.split("|", 1)[0] for m in _WIKILINK_RE.findall(text)}
        assert others & linked, f"{f.name} 未链接同 cluster 论文"


def test_moc_groups_by_cluster(vault_dir, papers, analyses):
    moc = next(f for f in os.listdir(vault_dir) if f.startswith("00-MOC-"))
    text = (vault_dir / moc).read_text(encoding="utf-8")
    for cluster in {a.cluster for a in analyses.values()}:
        assert f"## {cluster}" in text, f"MOC 缺少 cluster 分组 {cluster}"
    # MOC 应列出所有论文笔记
    note_stems = {f.stem for f in (vault_dir / "papers").glob("*.md")}
    for stem in note_stems:
        assert f"[[{stem}]]" in text, f"MOC 未列出笔记 {stem}"
    # MOC frontmatter 也应是合法 YAML
    _parse_frontmatter(text)


# ---------------- Excel ----------------

@pytest.fixture()
def xlsx_path(tmp_path, papers, analyses):
    out = tmp_path / "检索报告.xlsx"
    write_excel(QUERY, papers, analyses, str(out))
    return out


def test_excel_structure(xlsx_path, papers):
    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == [SHEET_NAME]
    ws = wb[SHEET_NAME]
    assert ws.max_row == len(papers) + 1
    assert ws.max_column == len(HEADERS)
    headers = [ws.cell(row=1, column=c).value for c in range(1, len(HEADERS) + 1)]
    assert headers == HEADERS
    # 表头加粗 + 浅灰填充
    h = ws.cell(row=1, column=1)
    assert h.font.bold
    assert h.fill.start_color.rgb in ("00D9D9D9", "FFD9D9D9")
    # 冻结首行 + 自动筛选
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_excel_sorted_by_composite_desc(xlsx_path):
    wb = load_workbook(xlsx_path)
    ws = wb[SHEET_NAME]
    comp_col = HEADERS.index("综合分") + 1
    scores = [ws.cell(row=r, column=comp_col).value
              for r in range(2, ws.max_row + 1)]
    assert scores == sorted(scores, reverse=True), "应按综合分降序"
    # 序号从 1 连续编号
    idx_col = HEADERS.index("序号") + 1
    assert [ws.cell(row=r, column=idx_col).value
            for r in range(2, ws.max_row + 1)] == list(range(1, ws.max_row))


def test_excel_conditional_formatting(xlsx_path):
    wb = load_workbook(xlsx_path)
    ws = wb[SHEET_NAME]
    rules = list(ws.conditional_formatting)
    assert len(rules) >= 3, "相关度/权威度/综合分应有 3 色阶条件格式"
    sqrefs = " ".join(str(r.sqref) for r in rules)
    for header in ("相关度", "权威度", "综合分"):
        col = HEADERS.index(header) + 1
        letter = chr(ord("A") + col - 1)
        assert f"{letter}2" in sqrefs, f"缺少 {header} 列的条件格式"


def test_excel_pdf_status_colors(xlsx_path):
    wb = load_workbook(xlsx_path)
    ws = wb[SHEET_NAME]
    pdf_col = HEADERS.index("PDF状态") + 1
    fills = {}
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=pdf_col)
        fills[cell.value] = cell.fill.start_color.rgb
    assert fills["ok"].endswith("C6EFCE"), "ok 应为绿色"
    assert fills["paywalled"].endswith("FFEB9C"), "paywalled 应为黄色"
    assert fills["error"].endswith("FFC7CE"), "error 应为红色"


def test_excel_hyperlink_and_summary_column(xlsx_path):
    wb = load_workbook(xlsx_path)
    ws = wb[SHEET_NAME]
    link_col = HEADERS.index("链接") + 1
    hyperlinks = 0
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=link_col)
        if cell.hyperlink is not None:
            hyperlinks += 1
            assert cell.font.underline == "single"
            assert cell.font.color and cell.font.color.rgb.endswith("0563C1")
    assert hyperlinks == 5, "所有论文都有落地页链接"

    # 内容大意列宽 60 且自动换行
    summary_col = HEADERS.index("内容大意") + 1
    letter = chr(ord("A") + summary_col - 1)
    assert ws.column_dimensions[letter].width == 60
    for r in range(2, ws.max_row + 1):
        assert ws.cell(row=r, column=summary_col).alignment.wrap_text


def test_excel_authors_truncated(xlsx_path):
    wb = load_workbook(xlsx_path)
    ws = wb[SHEET_NAME]
    auth_col = HEADERS.index("作者") + 1
    authors = [ws.cell(row=r, column=auth_col).value
               for r in range(2, ws.max_row + 1)]
    assert any("et al." in a for a in authors), "超过 3 位作者应写 et al."
