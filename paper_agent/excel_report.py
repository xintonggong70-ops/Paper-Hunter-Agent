"""Excel 汇总报告生成（SPEC §12）。

write_excel() 生成单 sheet '检索结果'：
- 冻结首行、自动筛选、表头加粗浅灰填充
- 相关度/权威度/综合分 3 色阶条件格式
- PDF 状态单元格着色：ok 绿 / paywalled 黄 / error 红
- 链接列为超链接样式；内容大意列宽 60 且自动换行
- 按综合分（0.6*相关 + 0.4*权威）降序
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Paper, PaperAnalysis

SHEET_NAME = "检索结果"

HEADERS = [
    "序号", "标题", "作者", "年份", "期刊/会议", "来源", "DOI", "链接",
    "引用数", "相关度", "权威度", "综合分", "主题聚类", "PDF状态",
    "Zotero导入", "内容大意",
]

# 各列宽度（列名 -> 宽度）
COL_WIDTHS = {
    "序号": 6, "标题": 50, "作者": 28, "年份": 8, "期刊/会议": 22,
    "来源": 16, "DOI": 24, "链接": 32, "引用数": 8, "相关度": 9,
    "权威度": 9, "综合分": 9, "主题聚类": 18, "PDF状态": 12,
    "Zotero导入": 12, "内容大意": 60,
}

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

_LINK_FONT = Font(color="0563C1", underline="single")

_WRAP_ALIGN = Alignment(wrap_text=True, vertical="top")
_TOP_ALIGN = Alignment(vertical="top")

# PDF 状态着色（Excel 经典 好/中/差 配色）
_PDF_STYLES = {
    "ok": ("C6EFCE", "006100"),          # 绿
    "paywalled": ("FFEB9C", "9C6500"),   # 黄
    "error": ("FFC7CE", "9C0006"),       # 红
}


def composite_score(paper: Paper) -> float:
    return 0.6 * paper.relevance_score + 0.4 * paper.authority_score


def _authors_cell(authors: list[str]) -> str:
    if not authors:
        return ""
    if len(authors) > 3:
        return "; ".join(authors[:3]) + "; et al."
    return "; ".join(authors)


def _zotero_status(paper: Paper) -> str:
    return "已导入" if paper.zotero_key else "未导入"


def write_excel(query: str, papers: list[Paper],
                analyses: dict[str, PaperAnalysis], out_path: str) -> None:
    """写 Excel 汇总报告（SPEC §12）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    # 表头
    for col, header in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=header)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = _HEADER_ALIGN
        ws.column_dimensions[get_column_letter(col)].width = COL_WIDTHS[header]

    # 按综合分降序
    sorted_papers = sorted(papers, key=composite_score, reverse=True)

    col_idx = {h: i + 1 for i, h in enumerate(HEADERS)}

    for row, p in enumerate(sorted_papers, start=2):
        analysis = analyses.get(p.id)
        comp = composite_score(p)
        values = {
            "序号": row - 1,
            "标题": p.title,
            "作者": _authors_cell(p.authors),
            "年份": p.year if p.year is not None else "",
            "期刊/会议": p.venue or "",
            "来源": p.source,
            "DOI": p.doi or "",
            "引用数": p.citation_count if p.citation_count is not None else "",
            "相关度": round(p.relevance_score, 1),
            "权威度": round(p.authority_score, 1),
            "综合分": round(comp, 1),
            "主题聚类": analysis.cluster if analysis else "",
            "PDF状态": p.pdf_status,
            "Zotero导入": _zotero_status(p),
            "内容大意": analysis.summary if analysis else "",
        }
        for header, value in values.items():
            c = ws.cell(row=row, column=col_idx[header], value=value)
            c.alignment = _WRAP_ALIGN if header in ("标题", "内容大意") else _TOP_ALIGN

        # 链接列：超链接样式
        link_cell = ws.cell(row=row, column=col_idx["链接"], value=p.url or "")
        if p.url:
            link_cell.hyperlink = p.url
            link_cell.font = _LINK_FONT
        link_cell.alignment = _WRAP_ALIGN

        # PDF 状态着色
        pdf_cell = ws.cell(row=row, column=col_idx["PDF状态"])
        style = _PDF_STYLES.get(p.pdf_status)
        if style:
            fill_color, font_color = style
            pdf_cell.fill = PatternFill("solid", start_color=fill_color,
                                        end_color=fill_color)
            pdf_cell.font = Font(color=font_color)

    last_row = len(sorted_papers) + 1

    # 冻结首行 + 自动筛选
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last_row}"

    # 相关度/权威度/综合分：3 色阶条件格式
    for header in ("相关度", "权威度", "综合分"):
        letter = get_column_letter(col_idx[header])
        rule = ColorScaleRule(
            start_type="num", start_value=0, start_color="F8696B",
            mid_type="num", mid_value=50, mid_color="FFEB84",
            end_type="num", end_value=100, end_color="63BE7B",
        )
        ws.conditional_formatting.add(f"{letter}2:{letter}{last_row}", rule)

    wb.save(out_path)
